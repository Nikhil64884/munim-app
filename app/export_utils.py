"""
Munim Ji AI — Excel Export.

Generates a styled Excel report with:
  - Transactions sheet (all non-deleted transactions)
  - Customer Balances sheet (net owed per customer)
  - Auto-width columns, header fill, summary rows
  - Unguessable filename (uuid suffix)
  - Auto-cleanup of old files after EXPORT_RETENTION_HOURS
"""

import datetime
import hashlib
import logging
import os
import uuid

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import text

from app.config import config
from app.database import get_db
from app.money import paise_to_rupees

logger = logging.getLogger("MunimJi.export")

os.makedirs(config.EXPORTS_DIR, exist_ok=True)


def _cleanup_old_exports():
    """Delete export files older than config.EXPORT_RETENTION_HOURS."""
    cutoff = datetime.datetime.now().timestamp() - (config.EXPORT_RETENTION_HOURS * 3600)
    try:
        for fname in os.listdir(config.EXPORTS_DIR):
            # Only ever touch our own exports, never unrelated files.
            if not (fname.startswith("hisaab_") and fname.endswith(".xlsx")):
                continue
            fpath = os.path.join(config.EXPORTS_DIR, fname)
            if os.path.isfile(fpath) and os.path.getmtime(fpath) < cutoff:
                os.remove(fpath)
                logger.info(f"Removed expired export: {fname}")
    except Exception as e:
        logger.warning(f"Cleanup skipped: {e}")


def get_owner_export_token(owner_id: str) -> str:
    """Return an 8-character cryptographic hash binding an export file to its owner."""
    return hashlib.sha256(f"{owner_id}:{config.JWT_SECRET}".encode()).hexdigest()[:8]


def _safe_excel_text(value):
    """Prevent untrusted text from being interpreted as an Excel formula."""
    if isinstance(value, str) and value.lstrip()[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def generate_excel(owner_id: str) -> str | None:
    """
    Export one owner's transactions to a styled Excel file.

    Returns the filename (not full path) on success, None on failure.
    """
    _cleanup_old_exports()

    owner_token = get_owner_export_token(owner_id)
    timestamp = datetime.datetime.now().strftime("%Y-%m-%d_%H%M")
    token = uuid.uuid4().hex[:12]
    filename = f"hisaab_{owner_token}_{timestamp}_{token}.xlsx"
    filepath = os.path.join(config.EXPORTS_DIR, filename)

    try:
        wb = openpyxl.Workbook()

        # ---- Sheet 1: Transactions ----
        ws = wb.active
        ws.title = "Transactions"

        headers = ["Date", "Customer", "Item/Desc", "Amount (₹)", "Type", "Original Message"]
        ws.append(headers)

        header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        total_udhaar = 0
        total_jama = 0
        total_cash = 0

        with get_db() as db:
            rows = db.execute(
                text(
                    "SELECT date, customer_name, item_desc, amount_paise, transaction_type, "
                    "original_message FROM transactions "
                    "WHERE owner_id=:o AND is_deleted=0 ORDER BY date DESC"
                ),
                {"o": owner_id},
            ).fetchall()

            for r in rows:
                date_val, name, item, paise, tx_type, message = r
                amount_rupees = paise_to_rupees(paise or 0)
                ws.append([
                    date_val,
                    _safe_excel_text(name),
                    _safe_excel_text(item),
                    amount_rupees,
                    _safe_excel_text(tx_type),
                    _safe_excel_text(message),
                ])

                t = (tx_type or "").lower()
                if t in ("debit", "sale"):
                    total_udhaar += paise or 0
                elif t == "credit":
                    total_jama += paise or 0
                elif t == "cashsale":
                    total_cash += paise or 0

        # ---- Summary rows ----
        ws.append([])
        summary_font = Font(bold=True, color="1F4E79")
        s = ws.max_row + 1
        ws.cell(row=s, column=3, value="Total Udhaar").font = summary_font
        ws.cell(row=s, column=4, value=paise_to_rupees(total_udhaar)).font = summary_font

        s += 1
        ws.cell(row=s, column=3, value="Total Jama").font = summary_font
        ws.cell(row=s, column=4, value=paise_to_rupees(total_jama)).font = summary_font

        s += 1
        ws.cell(row=s, column=3, value="Total Cash Sale").font = summary_font
        ws.cell(row=s, column=4, value=paise_to_rupees(total_cash)).font = summary_font

        s += 1
        net = total_udhaar - total_jama
        ws.cell(row=s, column=3, value="Net Outstanding").font = Font(bold=True, color="CC0000")
        ws.cell(row=s, column=4, value=paise_to_rupees(net)).font = Font(bold=True, color="CC0000")

        # ---- Auto-width ----
        for col in ws.columns:
            max_len = 0
            letter = col[0].column_letter
            for cell in col:
                v = cell.value
                max_len = max(max_len, len(str(v)) if v is not None else 0)
            ws.column_dimensions[letter].width = min(max_len + 4, 50)

        # ---- Sheet 2: Customer Balances ----
        ws2 = wb.create_sheet("Customer Balances")
        ws2.append(["Customer", "Balance (₹)"])
        for cell in ws2[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        with get_db() as db:
            balances = db.execute(
                text(
                    """
                    SELECT c.name,
                      COALESCE(SUM(CASE WHEN t.transaction_type IN ('Debit','Sale')
                                        THEN t.amount_paise ELSE 0 END), 0)
                      -
                      COALESCE(SUM(CASE WHEN t.transaction_type='Credit'
                                        THEN t.amount_paise ELSE 0 END), 0)
                      AS bal
                    FROM customers c
                    LEFT JOIN transactions t
                      ON t.customer_id=c.id AND t.is_deleted=0
                    WHERE c.owner_id=:o
                    GROUP BY c.id, c.name
                    ORDER BY bal DESC
                    """
                ),
                {"o": owner_id},
            ).fetchall()

            for name, bal_paise in balances:
                ws2.append([_safe_excel_text(name), paise_to_rupees(bal_paise or 0)])

        for col in ws2.columns:
            max_len = 0
            letter = col[0].column_letter
            for cell in col:
                v = cell.value
                max_len = max(max_len, len(str(v)) if v is not None else 0)
            ws2.column_dimensions[letter].width = min(max_len + 4, 40)

        wb.save(filepath)
        logger.info(f"Excel generated: {filepath}")
        return filename

    except Exception as e:
        logger.error(f"Excel error: {e}", exc_info=True)
        return None
